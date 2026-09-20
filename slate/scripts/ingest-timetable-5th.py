"""
Same as ingest-timetable.py but for the BTech5thSem sheet -- this batch's
own actual semester, per the roll-number ranges given directly by the
team (see NOTES.md). Column layout differs from the 3rd-sem sheet (no
merged-header offset here), so the time-column mapping is separate.
"""
import openpyxl
import re
import uuid
import boto3

SOURCE = '/home/kavyan2/Downloads/TimeTable_IT_July-Dec26.xlsx'
TABLE_NAME = 'TimetableSlot-r6zl3l4akngx5g3dhqm7v5pzd4-NONE'
REGION = 'ap-south-1'

TIME_COLS = {
    2: ('08:00', '09:00'),
    3: ('09:00', '10:00'),
    4: ('10:00', '11:00'),
    5: ('11:00', '12:00'),
    6: ('12:00', '13:00'),
    8: ('14:30', '15:30'),
    9: ('15:30', '16:30'),
    10: ('16:30', '17:30'),
    11: ('17:30', '18:30'),
}
DAYS = {'MON', 'TUE', 'WED', 'THU', 'FRI'}
ENTRY_RE = re.compile(
    r'([A-Za-z.]+)\s*\(([A-Za-z]+)\)\s*-\s*Sec\s*([A-Za-z0-9,\s]+?)\s*\(([^)]+)\)'
)
FACULTY_ROW_RE = re.compile(r'([^,()]+?)\s*\(([^)]+)\)')


def parse_faculty_legend(ws, code_col=3, faculty_col=7, start_row=15, end_row=29):
    result = {}
    for r in range(start_row, end_row):
        code = ws.cell(row=r, column=code_col).value
        faculty_text = ws.cell(row=r, column=faculty_col).value
        if not code or not faculty_text:
            continue
        by_section = {}
        for name, secs in FACULTY_ROW_RE.findall(faculty_text):
            for sec in re.split(r'[,\s]+', secs.strip()):
                if re.fullmatch(r'[A-Za-z][0-9]?', sec):
                    by_section[sec] = name.strip()
        if by_section:
            result[code.strip()] = by_section
    return result


def extract(ws, max_row=13):
    day_rows = []
    last_day, last_start = None, None
    for r in range(3, max_row):
        a = ws.cell(row=r, column=1).value
        if a in DAYS:
            if last_day:
                day_rows.append((last_day, last_start, r - 1))
            last_day, last_start = a, r
    if last_day:
        day_rows.append((last_day, last_start, max_row - 1))

    rows = []
    for day, r0, r1 in day_rows:
        for col, (start, end) in TIME_COLS.items():
            texts = [ws.cell(row=r, column=col).value for r in range(r0, r1 + 1)]
            texts = [t for t in texts if t]
            if not texts:
                continue
            for line in '\n'.join(texts).split('\n'):
                line = line.strip()
                if not line:
                    continue
                m = ENTRY_RE.search(line)
                if not m:
                    continue
                code, _kind, secs, room = m.groups()
                for sec in re.split(r'[,\s]+', secs.strip()):
                    if sec:
                        rows.append({
                            'day': day, 'startTime': start, 'endTime': end,
                            'courseId': code, 'section': sec, 'room': room.strip(),
                        })
    return rows


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb['BTech5thSem']
    entries = extract(ws)
    faculty_lookup = parse_faculty_legend(ws)

    ddb = boto3.resource('dynamodb', region_name=REGION)
    table = ddb.Table(TABLE_NAME)

    written = 0
    with table.batch_writer() as batch:
        for e in entries:
            item = {
                'id': str(uuid.uuid4()),
                '__typename': 'TimetableSlot',
                'program': 'BTech',
                'branch': 'IT',
                'semester': 5,
                'section': e['section'],
                'day': e['day'],
                'startTime': e['startTime'],
                'endTime': e['endTime'],
                'courseId': e['courseId'],
                'room': e['room'],
                'createdAt': '2026-09-18T00:00:00.000Z',
                'updatedAt': '2026-09-18T00:00:00.000Z',
            }
            faculty = faculty_lookup.get(e['courseId'], {}).get(e['section'])
            if faculty:
                item['faculty'] = faculty
            batch.put_item(Item=item)
            written += 1

    print(f'Wrote {written} TimetableSlot rows (5th sem) to {TABLE_NAME}')


if __name__ == '__main__':
    main()
