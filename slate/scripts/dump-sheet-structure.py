"""
Dumps a sheet's complete raw structure -- every non-empty cell's value,
row, column, and merge range (if any) -- with zero interpretation. This
is the input to the Bedrock normalization step, not a parser itself:
the goal is to give the LLM everything a human looking at the actual
spreadsheet would see, including exactly which cells are merged, so it
can correctly determine real class durations instead of us guessing
column-by-column.
"""
import json
import sys
import openpyxl

SOURCE = '/home/kavyan2/Downloads/TimeTable_IT_July-Dec26.xlsx'


def col_letter(c):
    return openpyxl.utils.get_column_letter(c)


def dump_sheet(ws):
    merged = list(ws.merged_cells.ranges)

    def merge_for(coord):
        for mc in merged:
            if coord in mc:
                return str(mc)
        return None

    cells = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip():
                cells.append({
                    'coord': cell.coordinate,
                    'row': cell.row,
                    'col': cell.column,
                    'colLetter': col_letter(cell.column),
                    'value': str(cell.value),
                    'merge': merge_for(cell.coordinate),
                })

    return {
        'mergedRanges': [str(m) for m in merged],
        'cells': cells,
    }


def main():
    sheet_name = sys.argv[1] if len(sys.argv) > 1 else 'BTech5thSem'
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb[sheet_name]
    result = dump_sheet(ws)
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
