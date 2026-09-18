"""
Backfills TimetableSlot.sessionType (L/P/T) from the same real source
cells already used for extraction -- this was captured by the original
regex match but discarded (courseId, "kind", section, room = groups()
and "kind" was never stored). Re-reads the same two sheets, matches
each real entry against its existing DB row by courseId/day/section/
room/startTime, and sets sessionType.
"""
import openpyxl
import re
import boto3

SOURCE = '/home/kavyan2/Downloads/TimeTable_IT_July-Dec26.xlsx'
TABLE_NAME = 'TimetableSlot-r6zl3l4akngx5g3dhqm7v5pzd4-NONE'
REGION = 'ap-south-1'

TIME_COLS_3RD = {
    2: ('08:00', '09:00'), 3: ('09:00', '10:00'), 4: ('10:00', '11:00'),
    5: ('11:00', '12:00'), 7: ('12:00', '13:00'), 9: ('14:30', '15:30'),
    10: ('15:30', '16:30'), 11: ('16:30', '17:30'), 12: ('17:30', '18:30'),
}
TIME_COLS_5TH = {
    2: ('08:00', '09:00'), 3: ('09:00', '10:00'), 4: ('10:00', '11:00'),
    5: ('11:00', '12:00'), 6: ('12:00', '13:00'), 8: ('14:30', '15:30'),
    9: ('15:30', '16:30'), 10: ('16:30', '17:30'), 11: ('17:30', '18:30'),
}
DAYS = {'MON', 'TUE', 'WED', 'THU', 'FRI'}
ENTRY_RE = re.compile(
    r'([A-Za-z.]+)\s*\(([A-Za-z]+)\)\s*-\s*Sec\s*([A-Za-z0-9,\s]+?)\s*\(([^)]+)\)'
)


def extract(ws, time_cols, max_row):
    day_rows = []
    last_day, last_start = None, None
    for r in range(3, max_row):
        a = ws.cell(row=r, column=1).value
        if a in DAYS:
            if last_day:
                day_rows.append((last_day, last_start, r - 1))
            last_day, last_start = a, r
    if last_day:
        day_rows.append((last_day, last_start, max_row - 1))

    rows = []
    for day, r0, r1 in day_rows:
        for col, (start, _end) in time_cols.items():
            texts = [ws.cell(row=r, column=col).value for r in range(r0, r1 + 1)]
            texts = [t for t in texts if t]
            if not texts:
                continue
            for line in '\n'.join(texts).split('\n'):
                m = ENTRY_RE.search(line.strip())
                if not m:
                    continue
                code, kind, secs, room = m.groups()
                for sec in re.split(r'[,\s]+', secs.strip()):
                    if sec:
                        rows.append({
                            'day': day, 'startTime': start, 'courseId': code,
                            'section': sec, 'room': room.strip(), 'sessionType': kind.upper(),
                        })
    return rows


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    entries = (
        [(3, e) for e in extract(wb['BTech3rdSem'], TIME_COLS_3RD, 40)]
        + [(5, e) for e in extract(wb['BTech5thSem'], TIME_COLS_5TH, 13)]
    )

    ddb = boto3.resource('dynamodb', region_name=REGION)
    table = ddb.Table(TABLE_NAME)
    items = table.scan()['Items']

    updated, not_found = 0, []
    for semester, e in entries:
        matches = [
            it for it in items
            if it.get('semester') == semester
            and it.get('day') == e['day']
            and it.get('courseId') == e['courseId']
            and it.get('section') == e['section']
            and it.get('room') == e['room']
            and it.get('startTime') == e['startTime']
        ]
        if len(matches) != 1:
            not_found.append((semester, e, len(matches)))
            continue
        table.update_item(
            Key={'id': matches[0]['id']},
            UpdateExpression='SET sessionType = :s',
            ExpressionAttributeValues={':s': e['sessionType']},
        )
        updated += 1

    print(f'updated {updated} rows')
    print(f'not matched: {len(not_found)}')
    for nf in not_found:
        print(nf)


if __name__ == '__main__':
    main()
