"""
Timetable reader + validator (ingestion pipeline steps 1-3).

Offline/original version. The deployed reader is
amplify/functions/parse-timetable/reader.ts, which also handles sheets
whose classes name no section (single-section batches, e.g. ECE).

  1. read_sheet():   spreadsheet -> neutral cells with real geometry. Hour
                     columns come from the header row's own merges, day
                     blocks from column A's merges -- no hardcoded column
                     numbers, so a differently laid-out sheet still reads.
  2. map_entries():  cell text -> TimetableSlot rows. Deterministic regex
                     for now; this is the step Bedrock replaces once the
                     account has model access (same input, same output).
  3. validate():     checks the rows against the sheet's own course legend
                     (codes, L-T-P-S weekly hours) plus room/section clashes.

Duration rule: a merge covering several hours for the FULL height of the
day block is a multi-hour class. A merge covering several hours on only
PART of the block (one sub-row) is ambiguous -- it can be a real 2-hour
class sitting above a 1-hour one, or just a wider text box. Those are
resolved per entry by which reading makes each section's weekly hours
match the legend's L-T-P-S; anything the legend can't settle is flagged.

Usage:
  python3 scripts/timetable_reader.py <file.xlsx> <sheet> [--out rows.json]
"""
import argparse
import json
import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.utils import get_column_letter

DAYS = ['MON', 'TUE', 'WED', 'THU', 'FRI', 'SAT']
TIME_RANGE_RE = re.compile(r'^\s*(\d{1,2})[:.](\d{2})\s*-\s*(\d{1,2})[:.](\d{2})\s*$')
ENTRY_RE = re.compile(
    r'^\s*([A-Za-z][A-Za-z.&]*)\s*\(\s*([LTP])\s*\)\s*-\s*(?:Sec\s*)?(.+?)\s*\(([^)]+)\)\s*$'
)
SECTION_RE = re.compile(r'^[A-Z]\d?$')


def to_24h(h, m):
    h = int(h)
    if h < 8:  # sheet writes afternoon hours as 1:00, 2:30, ...
        h += 12
    return f'{h:02d}:{m}'


def merge_of(ws, row, col):
    for m in ws.merged_cells.ranges:
        if m.min_row <= row <= m.max_row and m.min_col <= col <= m.max_col:
            return m
    return None


# ---------------------------------------------------------------- step 1

def read_sheet(ws):
    title = str(ws.cell(row=1, column=1).value or '')

    header_row = None
    for r in range(1, 15):
        hits = [c for c in range(1, ws.max_column + 1)
                if TIME_RANGE_RE.match(str(ws.cell(row=r, column=c).value or ''))]
        if len(hits) >= 3:
            header_row = r
            break
    if header_row is None:
        raise ValueError('no header row with time ranges found')

    hours = []  # [{start, end, cols}]
    for c in range(1, ws.max_column + 1):
        m = TIME_RANGE_RE.match(str(ws.cell(row=header_row, column=c).value or ''))
        if not m:
            continue
        start, end = to_24h(m[1], m[2]), to_24h(m[3], m[4])
        if start == '13:00':  # lunch column
            continue
        mr = merge_of(ws, header_row, c)
        cols = list(range(mr.min_col, mr.max_col + 1)) if mr else [c]
        hours.append({'start': start, 'end': end, 'cols': cols})
    col_to_hour = {c: i for i, h in enumerate(hours) for c in h['cols']}

    blocks = []  # [{day, rows}]
    r = header_row + 1
    while r <= ws.max_row:
        v = str(ws.cell(row=r, column=1).value or '').strip().upper()
        if v in DAYS:
            mr = merge_of(ws, r, 1)
            last = mr.max_row if mr else r
            blocks.append({'day': v, 'rows': list(range(r, last + 1))})
            r = last + 1
        elif blocks:
            break
        else:
            r += 1

    cells = []
    for b in blocks:
        for row in b['rows']:
            for c in sorted(col_to_hour):
                v = ws.cell(row=row, column=c).value
                if v is None or not str(v).strip() or str(v).strip().upper() == 'LUNCH':
                    continue
                mr = merge_of(ws, row, c)
                cols = range(mr.min_col, mr.max_col + 1) if mr else [c]
                rows = set(range(mr.min_row, mr.max_row + 1)) if mr else {row}
                hour_idx = sorted({col_to_hour[x] for x in cols if x in col_to_hour})
                cells.append({
                    'coord': f'{get_column_letter(c)}{row}',
                    'day': b['day'],
                    'hours': hour_idx,
                    'fullHeight': set(b['rows']) <= rows,
                    'text': str(v),
                })

    return {
        'title': title,
        'hours': hours,
        'cells': cells,
        'legend': read_legend(ws, blocks[-1]['rows'][-1] + 1 if blocks else header_row + 1),
    }


def read_legend(ws, from_row):
    for r in range(from_row, ws.max_row + 1):
        headers = {str(ws.cell(row=r, column=c).value or '').strip().lower(): c
                   for c in range(1, ws.max_column + 1)}
        if 'course code' in headers:
            break
    else:
        return {}

    def col(*names):
        for k, c in headers.items():
            if any(n in k for n in names):
                return c
        return None

    code_c, ltps_c, name_c, fac_c = col('course code'), col('l-t-p'), col('course name'), col('facult')
    legend = {}
    for rr in range(r + 1, ws.max_row + 1):
        code = str(ws.cell(row=rr, column=code_c).value or '').strip()
        if not code:
            if legend:
                break
            continue
        ltps_raw = str(ws.cell(row=rr, column=ltps_c).value or '') if ltps_c else ''
        nums = [int(x) for x in re.split(r'\s*[-–—]\s*', ltps_raw.strip()) if x.isdigit()]
        legend[code] = {
            'name': str(ws.cell(row=rr, column=name_c).value or '').strip() if name_c else '',
            'ltps': nums if len(nums) == 4 else None,
            'faculty': parse_faculty(str(ws.cell(row=rr, column=fac_c).value or '') if fac_c else ''),
        }
    return legend


def parse_faculty(text):
    """'Dr. X (A, B1), Dr. Y (B2, C)' -> {A: X, B1: X, ...}; a bare name -> {'*': name}."""
    text = text.strip()
    if not text:
        return {}
    by_sec = {}
    for name, secs in re.findall(r'([^,()]+?)\s*\(([^)]+)\)', text):
        tokens = [s for s in re.split(r'[,\s]+', secs.strip()) if s]
        if tokens and all(SECTION_RE.match(s) for s in tokens):
            for s in tokens:
                by_sec[s] = name.strip()
    return by_sec or {'*': text}


# ---------------------------------------------------------------- step 2

def map_entries(sheet):
    hours = sheet['hours']
    rows, skipped = [], []
    for cell in sheet['cells']:
        long_ = len(cell['hours']) > 1
        ambiguous = long_ and not cell['fullHeight']
        for line in cell['text'].split('\n'):
            line = line.strip()
            if not line:
                continue
            m = ENTRY_RE.match(line)
            if not m:
                skipped.append({'coord': cell['coord'], 'day': cell['day'], 'text': line,
                                'reason': 'not in "CODE (L/T/P) - Sec X (ROOM)" form'})
                continue
            code, kind, secs, room = m.groups()
            sections = [s for s in re.split(r'[,\s]+', secs.strip()) if s]
            if not all(SECTION_RE.match(s) for s in sections):
                skipped.append({'coord': cell['coord'], 'day': cell['day'], 'text': line,
                                'reason': f'no section, only a cohort label ("{secs.strip()}")'})
                continue
            first, last = cell['hours'][0], cell['hours'][-1]
            for sec in sections:
                rows.append({
                    'day': cell['day'],
                    'startTime': hours[first]['start'],
                    'endTime': hours[last if long_ and not ambiguous else first]['end'],
                    'longEndTime': hours[last]['end'] if ambiguous else None,
                    'courseId': code.strip(),
                    'sessionType': kind.upper(),
                    'section': sec,
                    'room': re.sub(r'\s+', '', room),
                    'source': cell['coord'],
                    'duration': 'ambiguous' if ambiguous else ('merge' if long_ else 'single'),
                })
    return rows, skipped


# ---------------------------------------------------------------- step 3

def minutes(t):
    h, m = t.split(':')
    return int(h) * 60 + int(m)


def atomic_groups(rows):
    """Each real student group: B1/B2 if the sheet splits B, else B."""
    secs = {r['section'] for r in rows}
    groups = set()
    for s in secs:
        if len(s) == 1 and any(len(x) == 2 and x[0] == s for x in secs):
            continue
        groups.add(s)
    return sorted(groups)


def applies(row_section, group):
    return row_section == group or row_section == group[0]


def weekly_hours(rows, group, code):
    lt = p = 0.0
    for r in rows:
        if r['courseId'] == code and applies(r['section'], group):
            h = (minutes(r['endTime']) - minutes(r['startTime'])) / 60
            if r['sessionType'] == 'P':
                p += h
            else:
                lt += h
    return lt, p


def validate(rows, legend):
    issues = []
    groups = atomic_groups(rows)

    for code in sorted({r['courseId'] for r in rows} - set(legend)):
        issues.append({'type': 'unknown-course', 'course': code,
                       'detail': 'course code not in the sheet\'s course legend'})

    # Resolve ambiguous merges against L-T-P-S: take the long reading only
    # where every section it covers is short of its legend hours by at
    # least that much.
    deficit = {}
    for code, info in legend.items():
        if not info['ltps']:
            continue
        l, t, p, _ = info['ltps']
        for g in groups:
            have_lt, have_p = weekly_hours(rows, g, code)
            if have_lt or have_p:
                deficit[(code, g, 'LT')] = (l + t) - have_lt
                deficit[(code, g, 'P')] = p - have_p

    for r in rows:
        if r['duration'] != 'ambiguous':
            continue
        extra = (minutes(r['longEndTime']) - minutes(r['endTime'])) / 60
        kind = 'P' if r['sessionType'] == 'P' else 'LT'
        affected = [g for g in groups if applies(r['section'], g)]
        keys = [(r['courseId'], g, kind) for g in affected]
        if keys and all(k in deficit and deficit[k] >= extra for k in keys):
            r['endTime'] = r['longEndTime']
            for k in keys:
                deficit[k] -= extra
            r['duration'] = 'ambiguous->long (matches L-T-P-S)'
        elif keys and all(k in deficit for k in keys):
            r['duration'] = 'ambiguous->short (matches L-T-P-S)'
        else:
            r['duration'] = 'ambiguous->short (UNVERIFIED, no L-T-P-S)'
            issues.append({'type': 'unverified-duration', 'row': short(r),
                           'detail': 'merge spans several hours on one sub-row; no legend hours to decide 1hr vs longer'})

    for (code, g, kind), d in sorted(deficit.items()):
        if d != 0:
            issues.append({'type': 'hours-mismatch', 'course': code, 'section': g, 'kind': kind,
                           'detail': f'legend says {d:+g}h/week vs what the sheet shows'})

    by_day = defaultdict(list)
    for r in rows:
        by_day[r['day']].append(r)
    for day, rs in by_day.items():
        for i, a in enumerate(rs):
            for b in rs[i + 1:]:
                if not (a['startTime'] < b['endTime'] and b['startTime'] < a['endTime']):
                    continue
                if a['room'] == b['room'] and a['courseId'] != b['courseId']:
                    issues.append({'type': 'room-clash', 'rows': [short(a), short(b)]})
                if a['courseId'] != b['courseId'] and any(
                        applies(a['section'], g) and applies(b['section'], g) for g in groups):
                    issues.append({'type': 'section-clash', 'rows': [short(a), short(b)]})
    return issues


def short(r):
    return f"{r['courseId']} ({r['sessionType']}) Sec {r['section']} {r['day']} {r['startTime']}-{r['endTime']} [{r['source']}]"


# ---------------------------------------------------------------- main

def batch_from_title(title):
    sem = re.search(r'(\d+)\s*(?:st|nd|rd|th)\s+Semester', title, re.I)
    branch = re.search(r'B\.?\s*Tech\.?\s*\((\w+)\)', title, re.I)
    return {
        'program': 'BTech' if re.search(r'B\.?\s*Tech', title, re.I) else None,
        'branch': branch[1] if branch else None,
        'semester': int(sem[1]) if sem else None,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('file')
    ap.add_argument('sheet')
    ap.add_argument('--out')
    args = ap.parse_args()

    ws = openpyxl.load_workbook(args.file, data_only=True)[args.sheet]
    sheet = read_sheet(ws)
    batch = batch_from_title(sheet['title'])
    rows, skipped = map_entries(sheet)
    issues = validate(rows, sheet['legend'])

    for r in rows:
        fac = sheet['legend'].get(r['courseId'], {}).get('faculty', {})
        r['faculty'] = fac.get(r['section']) or fac.get(r['section'][0]) or fac.get('*')
        r.update(batch)
        r.pop('longEndTime', None)

    print(f"{args.sheet}: {batch}")
    print(f"  {len(sheet['hours'])} hour columns, {len(sheet['cells'])} cells, "
          f"{len(sheet['legend'])} legend courses")
    print(f"  {len(rows)} rows, {len(skipped)} lines skipped, {len(issues)} issues")
    amb = [r for r in rows if r['duration'].startswith('ambiguous')]
    if amb:
        print('\n  ambiguous merges:')
        for r in amb:
            print(f"    {short(r)}  -> {r['duration']}")
    if skipped:
        print('\n  skipped lines:')
        for s in skipped:
            print(f"    [{s['coord']} {s['day']}] {s['text']!r}: {s['reason']}")
    if issues:
        print('\n  issues:')
        for i in issues:
            print('    ' + json.dumps(i))

    if args.out:
        with open(args.out, 'w') as f:
            json.dump({'batch': batch, 'rows': rows, 'skipped': skipped, 'issues': issues}, f, indent=2)
        print(f'\n  wrote {args.out}')


if __name__ == '__main__':
    sys.exit(main())
